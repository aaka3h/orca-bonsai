"""Local-only persistent chat/RAG regression fixtures; no personal files are read."""
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "outputs"))
import orca_memory as memory


class MemoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="memory-test-", dir=ROOT / "work")
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.db_path = self.directory / "private-memory" / "memory.sqlite3"
        self.store = memory.MemoryStore(self.db_path)

    def text(self, name, content, encoding="utf-8"):
        path = self.directory / name
        path.write_text(content, encoding=encoding)
        return path

    def test_restart_persists_chat_current_selection_titles_and_document_index(self):
        first = self.store.new_conversation()
        self.store.append_message(first, "user", "Project cobalt delivery date")
        self.store.append_message(first, "assistant", "The agreed date was Friday.")
        second = self.store.new_conversation()
        self.store.append_message(second, "user", "Independent other chat")
        self.store.set_current_conversation(first)
        path = self.text("delivery.md", "Cobalt shipment is scheduled for 19 October.")
        self.assertEqual(self.store.index_document(str(path))["status"], "indexed")
        restarted = memory.MemoryStore(self.db_path)
        self.assertEqual(restarted.current_conversation_id(), first)
        self.assertEqual(restarted.load_messages(first), [
            {"role": "user", "content": "Project cobalt delivery date"},
            {"role": "assistant", "content": "The agreed date was Friday."}])
        self.assertEqual(len(restarted.load_messages(second)), 1)
        titles = {row["id"]: row["title"] for row in restarted.list_conversations()}
        self.assertEqual(titles[first], "Project cobalt delivery date")
        self.assertEqual(restarted.search_documents("cobalt")[0]["name"], "delivery.md")
        self.assertEqual(len(restarted.list_documents()), 1)

    def test_conversation_isolation_explicit_recall_recent_order_and_delete(self):
        earlier = self.store.new_conversation()
        self.store.append_message(earlier, "user", "Our capybara nickname was Captain Cobalt")
        current = self.store.new_conversation()
        for number in range(7):
            self.store.append_message(current, "user" if number % 2 == 0 else "assistant", f"Current message {number}")
        context = self.store.context(current, "capybara")
        self.assertEqual(context["past_conversations"], [])
        self.assertEqual([m["content"] for m in context["recent_messages"]], [f"Current message {number}" for number in range(3, 7)])
        recalled = self.store.context(current, "What did I tell you earlier about capybara?")
        self.assertEqual(len(recalled["past_conversations"]), 1)
        self.assertIn("Captain Cobalt", recalled["past_conversations"][0]["content"])
        self.assertNotIn("Current message", recalled["past_conversations"][0]["content"])
        self.store.delete_conversation(earlier)
        self.assertEqual(self.store.context(current, "remember capybara")["past_conversations"], [])
        self.store.delete_conversation(current)
        self.assertEqual(self.store.load_messages(current), [])
        self.assertIsNone(self.store.current_conversation_id())
        with self.assertRaises(ValueError):
            self.store.append_message(current, "user", "No resurrected deleted chat")
        with self.assertRaises(ValueError):
            self.store.set_current_conversation(current)

    def test_explicit_recall_restores_older_fact_from_current_conversation(self):
        current = self.store.new_conversation()
        self.store.append_message(current, "user", "My terrapin nickname is Navigator Plum.")
        for number in range(9):
            self.store.append_message(current, "user" if number % 2 == 0 else "assistant", f"Other topic message {number}")
        self.assertNotIn("Navigator Plum", json.dumps(self.store.context(current, "terrapin")))
        recalled = self.store.context(current, "Remember my terrapin nickname?")
        self.assertEqual(len(recalled["past_conversations"]), 1)
        self.assertIn("Navigator Plum", recalled["past_conversations"][0]["content"])
        self.assertEqual(len(recalled["recent_messages"]), 4)
        self.assertLessEqual(len(json.dumps(recalled, ensure_ascii=False)), 2600)
        self.store.append_message(current, "user", "Recent terrapin detail.")
        recalled = self.store.context(current, "Remember my terrapin nickname?")
        self.assertNotIn("Recent terrapin detail.", [item["content"] for item in recalled["past_conversations"]])
        self.assertIn("Navigator Plum", json.dumps(recalled["past_conversations"]))

    def test_selected_and_nonselected_chat_deletion_clear_recall_and_selection(self):
        earlier = self.store.new_conversation()
        self.store.append_message(earlier, "user", "Older goldfinch conversation.")
        selected = self.store.new_conversation()
        self.store.append_message(selected, "user", "Selected puffin conversation.")
        self.store.delete_conversation(earlier)
        self.assertEqual(self.store.current_conversation_id(), selected)
        self.assertEqual(self.store.load_messages(earlier), [])
        self.assertEqual(self.store.context(selected, "remember goldfinch")["past_conversations"], [])
        with self.store._connect() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM message_search WHERE conversation_id=?", (earlier,)).fetchone()[0], 0)
        self.store.delete_conversation("missing-chat")
        self.assertEqual(self.store.current_conversation_id(), selected)
        self.store.delete_conversation(selected)
        reopened = memory.MemoryStore(self.db_path)
        self.assertIsNone(reopened.current_conversation_id())
        self.assertEqual(reopened.list_conversations(), [])
        self.assertEqual(reopened.context(None, "remember puffin")["past_conversations"], [])
        with reopened._connect() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM messages").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT count(*) FROM message_search").fetchone()[0], 0)
            self.assertIsNone(db.execute("SELECT value FROM settings WHERE key='current_chat'").fetchone())

    def test_clear_all_history_preserves_reference_documents_and_survives_restart(self):
        reference = self.text("reference.md", "Reference-only kingfisher calibration data.")
        self.assertEqual(self.store.index_document(str(reference))["status"], "indexed")
        source = self.store.search_documents("kingfisher")[0]
        identifiers = []
        for content in ("Remember chat-only grouse notes.", "Remember chat-only lapwing notes."):
            identifier = self.store.new_conversation()
            identifiers.append(identifier)
            self.store.append_message(identifier, "user", content)
            self.store.append_message(identifier, "assistant", "Acknowledged.")
        with self.store._connect() as db:
            db.execute("INSERT INTO settings VALUES ('unrelated_preference','keep')")
        self.assertEqual(self.store.clear_conversations(), 2)
        reopened = memory.MemoryStore(self.db_path)
        self.assertEqual(reopened.list_conversations(), [])
        self.assertIsNone(reopened.current_conversation_id())
        for identifier in identifiers:
            self.assertEqual(reopened.load_messages(identifier), [])
        self.assertEqual(reopened.context(None, "remember grouse lapwing")["past_conversations"], [])
        self.assertEqual(len(reopened.list_documents()), 1)
        self.assertEqual(reopened.search_documents("kingfisher")[0]["source_id"], source["source_id"])
        self.assertEqual(reference.read_text(), "Reference-only kingfisher calibration data.")
        with reopened._connect() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM messages").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT count(*) FROM message_search").fetchone()[0], 0)
            self.assertIsNone(db.execute("SELECT value FROM settings WHERE key='current_chat'").fetchone())
            self.assertEqual(db.execute("SELECT value FROM settings WHERE key='unrelated_preference'").fetchone()[0], "keep")
        self.assertEqual(reopened.clear_conversations(), 0)
        fresh = reopened.new_conversation()
        reopened.append_message(fresh, "user", "Fresh start")
        self.assertEqual(reopened.current_conversation_id(), fresh)
        self.assertEqual(len(reopened.load_messages(fresh)), 1)

    def test_clear_history_is_atomic_if_a_delete_fails(self):
        current = self.store.new_conversation()
        self.store.append_message(current, "user", "Atomic sandpiper facts.")
        with self.store._connect() as db:
            db.execute("CREATE TRIGGER block_chat_delete BEFORE DELETE ON conversations BEGIN SELECT RAISE(ABORT,'fixture failure'); END")
        with self.assertRaises(sqlite3.IntegrityError):
            self.store.clear_conversations()
        reopened = memory.MemoryStore(self.db_path)
        self.assertEqual(reopened.current_conversation_id(), current)
        self.assertEqual(reopened.load_messages(current)[0]["content"], "Atomic sandpiper facts.")
        with reopened._connect() as db:
            rows = db.execute('SELECT content FROM message_search WHERE message_search MATCH ?', ('"sandpiper"',)).fetchall()
            self.assertEqual(len(rows), 1)
            db.execute("DROP TRIGGER block_chat_delete")
        self.assertEqual(reopened.clear_conversations(), 1)

    def test_unchanged_reindex_removal_and_source_ids_are_stable_across_queries(self):
        path = self.text("plan.md", "Cobalt design contains alfalfa and dandelion.")
        self.assertEqual(self.store.index_document(str(path))["status"], "indexed")
        first = self.store.search_documents("alfalfa")[0]
        self.assertEqual(first["source_id"], self.store.search_documents("dandelion")[0]["source_id"])
        other = self.text("another.md", "A different alfalfa strategy.")
        self.store.index_document(str(other))
        for hit in self.store.search_documents("alfalfa"):
            if hit["path"] == str(path):
                self.assertEqual(hit["source_id"], first["source_id"])
            else:
                self.assertNotEqual(hit["source_id"], first["source_id"])
        with patch.object(memory.subprocess, "run", side_effect=AssertionError("unchanged text should not be extracted")):
            self.assertEqual(self.store.index_document(str(path))["status"], "unchanged")
        path.write_text("Replacement text contains kestrel and no old keywords.")
        self.assertEqual(self.store.index_document(str(path))["status"], "indexed")
        self.assertEqual(self.store.search_documents("dandelion"), [])
        replacement = self.store.search_documents("kestrel")[0]
        self.assertEqual(replacement["document_id"], first["document_id"])
        self.store.remove_document(replacement["document_id"])
        self.assertEqual(self.store.search_documents("kestrel"), [])
        self.assertTrue(path.exists(), "Removing indexed content must not delete the source file")
        self.assertEqual(len(self.store.list_documents()), 1)

    def test_literal_fts_queries_cannot_change_sql_or_activate_query_operators(self):
        path = self.text("notes.txt", "Only alfalfa is relevant.")
        self.store.index_document(str(path))
        self.store.index_document(str(self.text("other.txt", "Unrelated bumblebee.")))
        for query in ('"alfalfa" OR * )', 'alfalfa"; DROP TABLE documents; --', 'alfalfa NEAR(unknown)', '""" "alfalfa" [ ] : *'):
            with self.subTest(query=query):
                results = self.store.search_documents(query)
                self.assertTrue(results)
                self.assertEqual(results[0]["name"], "notes.txt")
        self.assertEqual(len(self.store.list_documents()), 2)
        self.assertEqual(self.store.search_documents("please summarize my documents"), [])
        self.assertEqual(self.store.search_documents("* \" : ()"), [])

    def test_actual_pdf_and_docx_extract_locations_and_limit_notes(self):
        from reportlab.pdfgen.canvas import Canvas
        from docx import Document
        pdf = self.directory / "manual.pdf"
        canvas = Canvas(str(pdf))
        canvas.drawString(70, 760, "Juniper maintenance happens on Tuesday.")
        canvas.showPage()
        canvas.drawString(70, 760, "Second page contains kingfisher controls.")
        canvas.save()
        indexed = self.store.index_document(str(pdf))
        self.assertEqual(indexed["status"], "indexed", indexed)
        hit = self.store.search_documents("kingfisher")[0]
        self.assertTrue(hit["location"].startswith("page 2"))
        docx = self.directory / "instructions.docx"
        document = Document()
        document.add_paragraph("Wisteria equipment requires weekly checks.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table orchid values are retained."
        document.sections[0].header.paragraphs[0].text = "Excludedheaderword"
        document.save(docx)
        indexed = self.store.index_document(str(docx))
        self.assertEqual(indexed["status"], "indexed", indexed)
        self.assertIn("body text only", " ".join(indexed["limits"]))
        self.assertTrue(self.store.search_documents("orchid")[0]["location"].startswith("document body"))
        self.assertEqual(self.store.search_documents("Excludedheaderword"), [])

    def test_unicode_and_utf16_and_multibyte_truncation(self):
        for name, encoding in (("日本語.txt", "utf-8"), ("utf16.txt", "utf-16")):
            path = self.text(name, "Café résumé 東京 नमस्ते. Orchid collection.", encoding=encoding)
            result = self.store.index_document(str(path))
            self.assertEqual(result["status"], "indexed", result)
        self.assertEqual(len(self.store.search_documents("東京")), 2)
        self.assertEqual(len(self.store.search_documents("café")), 2)
        # A valid UTF-8 file must not fail because the byte cap cuts a code point.
        long_text = self.text("long.txt", "A" + "🌿" * (memory.MAX_TEXT + 10))
        extracted = memory._extract(long_text)
        self.assertEqual(len(extracted["pieces"][0]["text"]), memory.MAX_TEXT)
        self.assertTrue(extracted["limits"])

    def test_context_citations_match_retrieved_source_and_serialized_bounds(self):
        chat = self.store.new_conversation()
        for number in range(12):
            self.store.append_message(chat, "user" if number % 2 else "assistant", "\x01" * 1000 + f"latest-{number}")
        for number in range(4):
            name = str(number) + ("quoted-\"\\" * 22) + ".txt"
            self.store.index_document(str(self.text(name, "Juniper evidence " + ("\\ \" unicode東京 " * 100))))
        retrieved = {item["source_id"]: item for item in self.store.search_documents("juniper")}
        context = self.store.context(chat, "juniper")
        self.assertLessEqual(len(json.dumps(context, ensure_ascii=False)), 2600)
        self.assertTrue(context["documents"])
        for source in context["documents"]:
            self.assertIn(source["source_id"], retrieved)
            hit = retrieved[source["source_id"]]
            self.assertEqual(source["name"], hit["name"])
            self.assertEqual(source["location"], hit["location"])
            self.assertTrue(hit["text"].startswith(source["text"].removesuffix(" … [excerpt]")))
        for budget in (2, 80, 150, 300, 500, 900, 1600, 2600):
            with self.subTest(budget=budget):
                context = self.store.context(chat, "juniper", limit=budget)
                self.assertLessEqual(len(json.dumps(context, ensure_ascii=False)), budget)
        self.assertLessEqual(len(memory._clip("very long content", 3)), 3)

    def test_explicit_chat_recall_is_not_crowded_out_by_document_matches(self):
        earlier = self.store.new_conversation()
        self.store.append_message(earlier, "user", "My favorite capybara has the nickname Admiral Indigo.")
        current = self.store.new_conversation()
        for index in range(4):
            self.store.index_document(str(self.text(f"reference-{index}.txt", "Capybara biology reference. " * 100)))
        context = self.store.context(current, "Remember what I told you about capybara earlier?")
        self.assertTrue(context["past_conversations"])
        self.assertIn("Admiral Indigo", context["past_conversations"][0]["content"])
        self.assertTrue(context["documents"])
        self.assertLessEqual(len(json.dumps(context, ensure_ascii=False)), 2600)

    def test_rejected_binary_missing_directory_device_unsupported_and_empty_leave_index(self):
        binary = self.directory / "binary.txt"
        binary.write_bytes(b"abc\0xyz")
        empty = self.text("empty.txt", "")
        unknown = self.text("program.py", "print('hello')")
        for value in (str(binary), str(self.directory / "missing.pdf"), str(self.directory), "/dev/null", str(unknown), str(empty), "https://example.test/file.pdf", None):
            with self.subTest(value=value):
                result = self.store.index_document(value)
                self.assertEqual(result["status"], "error")
                self.assertIn("error", result)
        self.assertEqual(self.store.list_documents(), [])
        oversized = self.directory / "oversized.txt"
        with oversized.open("wb") as file:
            file.truncate(memory.MAX_DOCUMENT_BYTES + 1)
        self.assertEqual(self.store.index_document(str(oversized))["status"], "error")
        self.assertEqual(list(self.db_path.parent.glob("extract-*")), [])

    def test_extraction_timeout_preserves_previous_index_and_cleans_private_snapshot(self):
        path = self.text("stable.txt", "Original salamander evidence")
        self.store.index_document(str(path))
        path.write_text("Changed woodland evidence")
        def fail(command, **kwargs):
            self.assertEqual(kwargs["timeout"], 35)
            snapshot = Path(command[-1])
            self.assertNotEqual(snapshot, path)
            self.assertEqual(snapshot.read_text(), "Changed woodland evidence")
            self.assertEqual(snapshot.stat().st_mode & 0o777, 0o600)
            self.assertEqual(snapshot.parent.stat().st_mode & 0o777, 0o700)
            self.assertNotIn("start_new_session", kwargs)
            raise subprocess.TimeoutExpired(command, 35)
        with patch.object(memory.subprocess, "run", side_effect=fail):
            result = self.store.index_document(str(path))
        self.assertEqual(result["status"], "error")
        self.assertIn("previous index was preserved", result["error"])
        self.assertTrue(self.store.search_documents("salamander"))
        self.assertEqual(self.store.search_documents("woodland"), [])
        self.assertEqual(list(self.db_path.parent.glob("extract-*")), [])

    def test_private_database_permissions_and_existing_parent_permissions_unchanged(self):
        self.assertEqual(self.db_path.stat().st_mode & 0o777, 0o600)
        self.assertEqual(self.db_path.parent.stat().st_mode & 0o777, 0o700)
        with self.store._connect() as db:
            db.execute("INSERT INTO settings VALUES ('test','yes')")
            for sidecar in self.db_path.parent.glob("memory.sqlite3-*"):
                self.assertEqual(sidecar.stat().st_mode & 0o777, 0o600)
        existing = self.directory / "existing-parent"
        existing.mkdir()
        existing.chmod(0o755)
        other = memory.MemoryStore(existing / "private.sqlite3")
        self.assertEqual(existing.stat().st_mode & 0o777, 0o755)
        self.assertEqual(other.path.stat().st_mode & 0o777, 0o600)

    def test_worker_resource_limits_are_applied_only_inside_worker(self):
        import resource
        with patch.object(resource, "getrlimit", return_value=(resource.RLIM_INFINITY, resource.RLIM_INFINITY)), patch.object(resource, "setrlimit") as set_limit:
            memory._bound_extraction_worker()
        self.assertIn(((resource.RLIMIT_AS, (1024 ** 3, 1024 ** 3)),), [(call.args,) for call in set_limit.call_args_list])
        self.assertIn(((resource.RLIMIT_CPU, (30, 30)),), [(call.args,) for call in set_limit.call_args_list])


if __name__ == "__main__":
    unittest.main()
